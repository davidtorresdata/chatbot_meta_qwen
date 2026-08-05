"""Knowledge ingestion.

Loads .txt / .md / .pdf / .xlsx files from a directory or path, chunks them
and stores embeddings in the local LanceDB table. This is the process used to
grow the knowledge base without touching code.

Excel files are converted to one text block per data row so each row is
searchable independently (see ``load_excel_rows``).
"""

from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path

from pypdf import PdfReader

from src.knowledge.chunker import TextChunker
from src.knowledge.embedding import EmbeddingProvider
from src.knowledge.vector_store import VectorStore

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".txt", ".md", ".markdown", ".pdf", ".xlsx", ".xlsm"}


@dataclass
class Document:
    text: str
    source: str


def load_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def load_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n\n".join(pages)


def load_excel_rows(path: Path) -> list[str]:
    """Convert an .xlsx workbook into one text block per data row.

    The first non-empty row of each sheet is treated as the column header
    (used to label every cell value). Each subsequent non-empty row becomes a
    single text block::

        [sheet 'Pricing', row 3] Product: Premium Widget | Price: 129.99 ...

    Rows that are entirely empty are skipped. Sheets without a header row will
    treat the first data row as the header — keep a header row in every sheet.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "Excel support requires openpyxl (pip install openpyxl)"
        ) from exc

    blocks: list[str] = []
    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            sheet_name = ws.title
            header: list[str] | None = None
            for row in ws.iter_rows():
                values = [cell.value for cell in row]
                if all(v is None or str(v).strip() == "" for v in values):
                    continue
                row_no = row[0].row
                if header is None:
                    header = [
                        str(v).strip() if v is not None and str(v).strip() else f"col{i + 1}"
                        for i, v in enumerate(values)
                    ]
                    continue
                cells = [
                    f"{header[i]}: {v}"
                    for i, v in enumerate(values)
                    if i < len(header) and v is not None and str(v).strip()
                ]
                if not cells:
                    continue
                blocks.append(f"[sheet '{sheet_name}', row {row_no}] " + " | ".join(cells))
    finally:
        wb.close()
    return blocks


def load_file_blocks(path: Path) -> list[str]:
    """Return the raw text blocks of a file.

    Text and PDF files produce a single block; Excel files produce one block
    per data row.
    """
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return [load_pdf(path)]
    if suffix in {".xlsx", ".xlsm"}:
        return load_excel_rows(path)
    return [load_text_file(path)]


def load_file(path: Path) -> str:
    """Load a whole file as a single string (legacy convenience helper)."""
    blocks = load_file_blocks(path)
    return "\n\n".join(blocks) if blocks else ""


def collect_files(path: Path) -> list[Path]:
    if path.is_file():
        files = [path]
    elif path.is_dir():
        files = [p for p in sorted(path.rglob("*")) if p.suffix.lower() in SUPPORTED_EXTENSIONS]
    else:
        raise FileNotFoundError(f"path does not exist: {path}")
    return files


def chunk_documents(files: list[Path], chunker: TextChunker) -> list[Document]:
    documents: list[Document] = []
    for file in files:
        try:
            blocks = load_file_blocks(file)
        except Exception as exc:  # keep going if one file is corrupt
            logger.warning("Could not read %s: %s", file, exc)
            continue
        index = 0
        for block in blocks:
            for chunk in chunker.split_text(block):
                if not chunk:
                    continue
                documents.append(
                    Document(text=chunk, source=f"{file}:{index}")
                )
                index += 1
    return documents


async def ingest_documents(
    vector_store: VectorStore,
    embedder: EmbeddingProvider,
    documents: list[Document],
    batch_size: int = 32,
) -> int:
    total = 0
    for start in range(0, len(documents), batch_size):
        batch = documents[start : start + batch_size]
        vectors = await embedder.embed([doc.text for doc in batch])
        vector_store.add(
            vectors=vectors,
            texts=[doc.text for doc in batch],
            metadatas=[{"source": doc.source} for doc in batch],
        )
        total += len(batch)
        logger.info("Ingested %d/%d chunks", total, len(documents))
    return total
